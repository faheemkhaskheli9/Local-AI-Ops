"""Turn files into plain text without an LLM.

Supported: .txt .md .html/.htm (stdlib), .pdf (pypdf), .docx (python-docx),
.csv/.json (stdlib). Install extras with: pip install -e ".[docs]"
"""
from __future__ import annotations

import csv
import io
import json
import re
from html.parser import HTMLParser
from pathlib import Path

SKIP_TAGS = {"script", "style", "noscript", "nav", "footer", "header", "svg", "form"}
BLOCK_TAGS = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6",
              "section", "article", "ul", "ol", "table", "pre", "blockquote"}


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in SKIP_TAGS:
            self.skip += 1
        elif tag in BLOCK_TAGS:
            self.parts.append("\n")
        if tag == "li" and not self.skip:
            self.parts.append("- ")
        if tag in {"h1", "h2", "h3"} and not self.skip:
            self.parts.append("#" * int(tag[1]) + " ")

    def handle_endtag(self, tag):
        if tag in SKIP_TAGS and self.skip:
            self.skip -= 1
        elif tag in BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def html_to_text(html: str) -> str:
    p = _HTMLText()
    p.feed(html)
    text = "".join(p.parts)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def pdf_to_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()


def docx_to_text(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    lines = []
    for para in d.paragraphs:
        style = (para.style.name or "").lower() if para.style is not None else ""
        txt = para.text.strip()
        if not txt:
            continue
        if style.startswith("heading"):
            level = re.sub(r"\D", "", style) or "1"
            lines.append("#" * int(level) + " " + txt)
        elif "list" in style:
            lines.append("- " + txt)
        else:
            lines.append(txt)
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n\n".join(lines)


def csv_to_text(text: str, max_rows: int = 50) -> str:
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return ""
    head, body = rows[0], rows[1:max_rows + 1]
    out = ["| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
    out += ["| " + " | ".join(r) + " |" for r in body]
    if len(rows) - 1 > max_rows:
        out.append(f"... {len(rows) - 1 - max_rows} more rows")
    return "\n".join(out)


def to_text(path: str | Path) -> str:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return pdf_to_text(path)
    if ext == ".docx":
        return docx_to_text(path)
    raw = path.read_text(encoding="utf-8", errors="replace")
    if ext in {".html", ".htm"}:
        return html_to_text(raw)
    if ext == ".csv":
        return csv_to_text(raw)
    if ext == ".json":
        return json.dumps(json.loads(raw), indent=1, ensure_ascii=False)
    return raw
